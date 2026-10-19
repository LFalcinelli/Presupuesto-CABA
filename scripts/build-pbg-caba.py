"""Build the denominator only. Usage: python scripts/build-pbg-caba.py --sources DIR.

Requires openpyxl and xlrd. DIR contains manifest.json and the five official
workbooks (annual, legacy, quarterNominal, quarterReal, rem). Raw files stay
outside public/ and outside the repository. No expenditure or price data change.
"""
import argparse
import hashlib
import json
from pathlib import Path

import xlrd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def build(source_dir):
    sources = json.loads((source_dir / 'manifest.json').read_text(encoding='utf-8'))
    for source in sources.values():
        assert hashlib.sha256((source_dir / source['file']).read_bytes()).hexdigest() == source['sha256']
    annual = openpyxl.load_workbook(source_dir / sources['annual']['file'], data_only=True).worksheets[0]
    legacy = xlrd.open_workbook(source_dir / sources['legacy']['file']).sheet_by_index(0)
    quarter = openpyxl.load_workbook(source_dir / sources['quarterNominal']['file'], data_only=True).worksheets[0]
    real = openpyxl.load_workbook(source_dir / sources['quarterReal']['file'], data_only=True).worksheets[0]
    rem = openpyxl.load_workbook(source_dir / sources['rem']['file'], data_only=True).worksheets[0]
    assert annual['A4'].value == quarter['A4'].value == real['A4'].value == 'Producto Geográfico Bruto'
    assert legacy.cell_value(4, 0) == 'Valor Agregado Bruto a precios básicos'
    assert rem['B104'].value == 'PIB a precios constantes' and rem['B110'].value == '2027' and rem['D105'].value == 'Mediana'
    old = {int(str(legacy.cell_value(1, c)).strip().rstrip('*').split('.')[0]): legacy.cell_value(4, c) * 1e6 for c in range(1, legacy.ncols)}
    observed = {int(str(annual.cell(3, c).value).rstrip('*')): annual.cell(4, c).value * 1e6 for c in range(2, 23)}
    ratios = {year: observed[year] / old[year] for year in range(2004, 2010)}
    link = ratios[2004]
    spread = max(ratios.values()) / min(ratios.values()) - 1
    quarterly2024 = [quarter.cell(4, c).value * 1e6 for c in range(2, 6)]
    quarterly2025 = [quarter.cell(4, c).value * 1e6 for c in range(8, 12)]
    discrepancy = sum(quarterly2024) / observed[2024] - 1
    scale = observed[2024] / sum(quarterly2024) if abs(discrepancy) > .005 else 1
    pbg2025 = sum(quarterly2025) * scale
    g2026 = sum(real.cell(4, c).value for c in [14, 15]) / sum(real.cell(4, c).value for c in [8, 9]) - 1
    g2027 = rem['D110'].value / 100
    history = json.loads((ROOT / 'data/history/execution-history.json').read_text(encoding='utf-8'))
    prices = {r['year']: r['priceIndexAnnual'] for r in history['rows'] if r['year'] >= 2025}
    pbg2026 = pbg2025 * (1 + g2026) * prices[2026] / prices[2025]
    pbg2027 = pbg2026 * (1 + g2027) * prices[2027] / prices[2026]
    nominalH1 = sum(quarter.cell(4, c).value for c in [14, 15]) / sum(quarter.cell(4, c).value for c in [8, 9]) - 1
    nominalEstimated = pbg2026 / pbg2025 - 1

    def provenance(key, cell, formula):
        s = sources[key]
        return dict(source={'page': s['page'], 'file': s['url']}, sourceSha256=s['sha256'], sourceSheet={'annual': annual.title, 'legacy': legacy.name, 'quarterNominal': quarter.title, 'quarterReal': real.title, 'rem': rem.title}[key], cell=cell, formula=formula)

    rows = []
    for year in range(1997, 2028):
        if year < 2004:
            if year not in old:
                continue
            row = dict(year=year, pbgNominal=old[year] * link, status='empalmado', **provenance('legacy', f'{openpyxl.utils.get_column_letter(year-1991)}5', 'VAB_base1993 * link.ratio2004'))
        elif year <= 2024:
            col = year - 2002
            provisional = '*' in str(annual.cell(3, col).value)
            row = dict(year=year, pbgNominal=observed[year], status='provisorio' if provisional else 'observado', **provenance('annual', f'{openpyxl.utils.get_column_letter(col)}4', 'published_millions * 1000000'))
        elif year == 2025:
            row = dict(year=year, pbgNominal=pbg2025, status='provisorio', **provenance('quarterNominal', 'H4:K4', 'sum(quarterly2025) * quarterlyControl.scale'))
        else:
            key = 'quarterReal' if year == 2026 else 'rem'
            row = dict(year=year, pbgNominal=pbg2026 if year == 2026 else pbg2027, status='estimado', **provenance(key, 'N4:O4 / H4:I4' if year == 2026 else 'D110', f'pbgNominal[{year-1}] * (1 + growth[{year}]) * priceIndexAnnual[{year}] / priceIndexAnnual[{year-1}]'))
        rows.append(row)
    return dict(updated='2026-10-19', metric='PBG de CABA a precios básicos, pesos corrientes', unit='pesos', numerator='execution-history.series.currentPrimary.rows[].nominal (referencia; no duplicado)', sources=sources,
        link=dict(anchorYear=2004, ratio2004=link, annualRatios=ratios, minimum=min(ratios.values()), maximum=max(ratios.values()), relativeRange=spread, exceedsThreePercent=spread > .03, anchorPublished=observed[2004], anchorLegacy=old[2004]),
        quarterlyControl=dict(annual2024=observed[2024], quarterly2024=quarterly2024, quarterly2025=quarterly2025, discrepancy=discrepancy, scale=scale, note='Los trimestres se publican a tasa anualizada: su promedio reproduce el anual. Se aplica el cociente anual 2024 / suma trimestral 2024, conforme al control autorizado.'),
        assumptions=dict(growth={'2026':g2026,'2027':g2027}, priceIndexAnnual=prices, priceAssumptions=history['priceAdjustment']['assumptions'], realGrowth2026Source=provenance('quarterReal','N4:O4 / H4:I4','sum(H1_2026) / sum(H1_2025) - 1'), realGrowth2027Source={**provenance('rem','D110','median_percent / 100'),'remMonth':'2026-09','published':'2026-10-06','scope':'PIB nacional como aproximación de CABA; pronóstico, no dato observado'}, nominalControl2026=dict(estimatedAnnualGrowth=nominalEstimated, observedH1Growth=nominalH1, differencePoints=(nominalEstimated-nominalH1)*100)), rows=rows)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--sources', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=ROOT/'data/history/pbg-caba.json')
    args = parser.parse_args()
    result = build(args.sources)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps({key:result[key] for key in ['link','quarterlyControl','assumptions']},ensure_ascii=False))
